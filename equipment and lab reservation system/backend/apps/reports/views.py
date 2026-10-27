import csv
from io import BytesIO

from django.db.models import Count, Sum
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.permissions import IsAdminUser
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.models import AuditLog
from apps.equipment.models import Equipment, EquipmentLoan
from apps.laboratories.models import Laboratory
from apps.reservations.models import Reservation


class DashboardSummaryView(APIView):
	permission_classes = [IsAdminUser]

	def get(self, request):
		equipment = Equipment.objects.filter(is_active=True)
		total_units = equipment.aggregate(total=Sum('quantity'))['total'] or 0
		loaned_units = EquipmentLoan.objects.filter(returned_at__isnull=True).aggregate(total=Sum('quantity'))['total'] or 0
		unavailable_units = equipment.filter(
			condition__in=[Equipment.Condition.DAMAGED, Equipment.Condition.MAINTENANCE]
		).aggregate(total=Sum('quantity'))['total'] or 0
		available_units = max(total_units - loaned_units - unavailable_units, 0)
		status_counts = dict(Reservation.objects.values_list('status').annotate(count=Count('id')))
		return Response({
			'equipment_count': equipment.count(),
			'equipment_units': total_units,
			'available_units': available_units,
			'unavailable_units': max(total_units - available_units, 0),
			'laboratory_count': Laboratory.objects.filter(is_active=True).count(),
			'reservation_count': Reservation.objects.count(),
			'pending_count': status_counts.get(Reservation.Status.PENDING, 0),
			'approved_count': status_counts.get(Reservation.Status.APPROVED, 0),
			'rejected_count': status_counts.get(Reservation.Status.DECLINED, 0),
			'cancelled_count': status_counts.get(Reservation.Status.CANCELLED, 0),
			'completed_count': status_counts.get(Reservation.Status.COMPLETED, 0),
			'overdue_loans': sum(loan.is_overdue for loan in EquipmentLoan.objects.filter(returned_at__isnull=True)),
			'recent_activity': list(AuditLog.objects.select_related('actor').values('action', 'summary', 'created_at', 'actor__username')[:8]),
		})


class ReservationReportView(APIView):
	permission_classes = [IsAdminUser]

	def get(self, request):
		start = parse_date(request.query_params.get('start', ''))
		end = parse_date(request.query_params.get('end', ''))
		reservations = Reservation.objects.select_related('equipment', 'laboratory').all()
		if start:
			reservations = reservations.filter(start_at__date__gte=start)
		if end:
			reservations = reservations.filter(start_at__date__lte=end)
		by_status = list(reservations.values('status').annotate(count=Count('id')).order_by('status'))
		top_equipment = list(
			EquipmentLoan.objects.values('equipment__name').annotate(borrowed=Sum('quantity')).order_by('-borrowed')[:10]
		)
		lab_usage = list(
			reservations.filter(laboratory__isnull=False).values('laboratory__name').annotate(bookings=Count('id')).order_by('-bookings')
		)
		reservation_rows = list(reservations.values_list(
			'id', 'requester__username', 'equipment__name', 'laboratory__name',
			'start_at', 'end_at', 'quantity', 'attendees', 'status', 'purpose',
		))
		overdue_loans = list(EquipmentLoan.objects.filter(returned_at__isnull=True).select_related('equipment', 'borrower'))
		damaged_equipment = list(Equipment.objects.filter(condition=Equipment.Condition.DAMAGED).values('name', 'asset_tag', 'quantity', 'location'))
		if request.query_params.get('export') == 'csv':
			response = HttpResponse(content_type='text/csv')
			response['Content-Disposition'] = 'attachment; filename="reservation-report.csv"'
			writer = csv.writer(response)
			writer.writerow(['ID', 'Requester', 'Equipment', 'Laboratory', 'Start', 'End', 'Quantity', 'Attendees', 'Status', 'Purpose'])
			writer.writerows(
				[
					value.isoformat() if hasattr(value, 'isoformat') else value
					for value in row
				]
				for row in reservation_rows
			)
			return response
		if request.query_params.get('export') == 'xlsx':
			workbook = Workbook()
			bookings = workbook.active
			bookings.title = 'Reservations'
			bookings.append(['ID', 'Requester', 'Equipment', 'Laboratory', 'Start', 'End', 'Quantity', 'Attendees', 'Status', 'Purpose'])
			for row in reservation_rows:
				bookings.append([
					timezone.localtime(value).replace(tzinfo=None) if hasattr(value, 'tzinfo') and value is not None else value
					for value in row
				])
			loan_sheet = workbook.create_sheet('Overdue and Unreturned')
			loan_sheet.append(['Equipment', 'Borrower', 'Quantity', 'Due', 'Overdue'])
			for loan in overdue_loans:
				loan_sheet.append([loan.equipment.name, loan.borrower.username, loan.quantity, timezone.localtime(loan.due_at).replace(tzinfo=None), loan.is_overdue])
			damaged_sheet = workbook.create_sheet('Damaged Equipment')
			damaged_sheet.append(['Name', 'Asset tag', 'Quantity', 'Location'])
			for row in damaged_equipment:
				damaged_sheet.append([row['name'], row['asset_tag'], row['quantity'], row['location']])
			output = BytesIO()
			workbook.save(output)
			response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
			response['Content-Disposition'] = 'attachment; filename="reservation-report.xlsx"'
			return response
		return Response({
			'by_status': by_status,
			'top_equipment': top_equipment,
			'laboratory_usage': lab_usage,
			'overdue_loans': [
				{'equipment': loan.equipment.name, 'borrower': loan.borrower.username, 'due_at': loan.due_at, 'overdue': loan.is_overdue}
				for loan in overdue_loans
			],
			'damaged_equipment': damaged_equipment,
			'total': reservations.count(),
		})
